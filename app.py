"""
app.py — Servidor Flask para Asistencia QR y Gestión Académica
Con persistencia en Cloudflare R2 (S3 compatible) para despliegue en Render (disco efímero).
- Los archivos se descargan a /tmp/ bajo demanda si el servidor se reinicia o duerme.
- Cada modificación se guarda localmente y se sube de inmediato a Cloudflare R2.
- Mantiene el 100% de la compatibilidad con openpyxl, colores, formatos y frontend original.
"""

import os
import sys
import io
import datetime
from flask import Flask, request, jsonify, send_file, send_from_directory
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import r2_storage

# Fix encoding para la consola
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

app = Flask(__name__, static_folder=None)

MATERIAS = {
    'optimizacion': {
        'nombre': 'Optimización',
        'horario': '9:00 - 11:00',
        'archivo': 'IIAT31-OPTIMIZACION-A1.xlsx'
    },
    'investigacion_operaciones': {
        'nombre': 'Investigación de Operaciones',
        'horario': '11:00 - 1:00',
        'archivo': 'IIAT33-INVESTIGACION DE OPERACIONES-A3.xlsx'
    },
    'teoria_decision': {
        'nombre': 'Teoría de la Decisión',
        'horario': '1:00 - 3:00',
        'archivo': 'IIAT32-TEORIA DE LA DECISION-A2.xlsx'
    }
}

# Estilos de celdas
FILL_GREEN = PatternFill('solid', fgColor='C6EFCE')
FONT_GREEN = Font(color='276221', bold=True)
FILL_RED = PatternFill('solid', fgColor='FFC7CE')
FONT_RED = Font(color='9C0006', bold=True)
FILL_HDR = PatternFill('solid', fgColor='1F4E79')
FONT_HDR = Font(bold=True, color='FFFFFF')
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')


def get_fecha_hoy_str():
    """Devuelve la fecha actual en formato DD/MM/YYYY."""
    return datetime.datetime.now().strftime('%d/%m/%Y')


def get_filepath(materia_key, force_download=False):
    """
    Obtiene la ruta local en /tmp/ del archivo de la materia.
    Asegura que esté descargado desde Cloudflare R2.
    """
    info = MATERIAS.get(materia_key)
    if not info:
        return None, None
    filename = info['archivo']
    local_path = r2_storage.ensure_file_local(filename, force_download=force_download)
    return local_path, filename


def guardar_workbook_seguro(wb, local_path, filename):
    """
    Guarda el workbook en el disco local (/tmp/) y lo sube inmediatamente a Cloudflare R2.
    """
    try:
        wb.save(local_path)
    except Exception as e:
        return False, f"Error al guardar localmente: {str(e)}"

    # Persistencia en la nube (Cloudflare R2)
    subido = r2_storage.upload_file_to_r2(filename)
    if not subido and r2_storage.is_r2_configured():
        return False, "Error al sincronizar el archivo con Cloudflare R2."

    return True, None


def get_target_sheet(wb, modo='asistencia'):
    """
    Retorna la hoja correspondiente según el modo:
    - modo 'noticia'   -> Hoja 'Noticia' (o 'Noticias')
    - modo 'asistencia' -> Hoja 'Hoja2' (o 'Asistencia')
    """
    m = str(modo or 'asistencia').lower().strip()
    if m == 'noticia':
        for sname in ['Noticia', 'Noticias']:
            if sname in wb.sheetnames:
                return wb[sname]
        # Si no existe, crear la hoja Noticia basada en Hoja2
        ws = wb.create_sheet(title='Noticia')
        ws_ref = wb['Hoja2'] if 'Hoja2' in wb.sheetnames else wb.worksheets[1]
        for c in range(1, ws_ref.max_column + 1):
            ws.cell(row=1, column=c, value=ws_ref.cell(row=1, column=c).value)
        for r in range(2, ws_ref.max_row + 1):
            ws.cell(row=r, column=1, value=ws_ref.cell(row=r, column=1).value)
            ws.cell(row=r, column=2, value=ws_ref.cell(row=r, column=2).value)
        return ws
    else:
        for sname in ['Hoja2', 'Asistencia']:
            if sname in wb.sheetnames:
                return wb[sname]
        return wb.worksheets[1] if len(wb.worksheets) > 1 else wb.active


def obtener_o_crear_columna_fecha(ws_target, fecha):
    """
    Busca la columna correspondiente a la fecha en la fila 1.
    Si no existe, la crea con el formato de encabezado correspondiente y ancho adecuado.
    """
    for col in range(3, ws_target.max_column + 1):
        if str(ws_target.cell(row=1, column=col).value or '').strip() == str(fecha).strip():
            return col

    col_fecha = ws_target.max_column + 1
    h_cell = ws_target.cell(row=1, column=col_fecha, value=str(fecha).strip())
    h_cell.fill = PatternFill('solid', fgColor='2E75B6')
    h_cell.font = FONT_HDR
    h_cell.alignment = ALIGN_CENTER
    ws_target.column_dimensions[get_column_letter(col_fecha)].width = 13
    return col_fecha


# ─── RUTAS PRINCIPALES ────────────────────────────────────────────────────────

@app.route('/')
def serve_index():
    """Sirve la interfaz web principal."""
    return send_from_directory('.', 'index.html')


@app.route('/api/materias', methods=['GET'])
def get_materias():
    """Devuelve la configuración de materias y horarios."""
    return jsonify({
        'ok': True,
        'materias': MATERIAS,
        'fecha_hoy': get_fecha_hoy_str()
    })


@app.route('/api/datos', methods=['GET'])
def get_datos_materia():
    """
    Carga los alumnos, fechas de clase y asistencias de ambas hojas (Asistencia y Noticia).
    Descarga previamente de R2 si el archivo no existe en /tmp/.
    """
    materia_key = request.args.get('materia', 'optimizacion')
    fecha_filtro = request.args.get('fecha', get_fecha_hoy_str())

    path, filename = get_filepath(materia_key)
    if not path or not os.path.exists(path):
        return jsonify({'ok': False, 'error': f"Archivo para '{materia_key}' no encontrado ni en local ni en R2."}), 404

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': f"Error al abrir Excel: {str(e)}"}), 500

    ws_asist = get_target_sheet(wb, 'asistencia')
    ws_notic = get_target_sheet(wb, 'noticia')
    
    headers = []
    col_fecha_asist = -1

    for col in range(1, ws_asist.max_column + 1):
        val = ws_asist.cell(row=1, column=col).value
        str_val = str(val).strip() if val is not None else ''
        headers.append(str_val)
        if col > 2 and str_val == fecha_filtro:
            col_fecha_asist = col

    col_fecha_notic = -1
    for col in range(1, ws_notic.max_column + 1):
        val = ws_notic.cell(row=1, column=col).value
        str_val = str(val).strip() if val is not None else ''
        if col > 2 and str_val == fecha_filtro:
            col_fecha_notic = col

    fechas_clase = [h for h in headers[2:] if h]

    # Mapeo rápido de noticia por número de cuenta
    noticia_map = {}
    if col_fecha_notic != -1:
        for r in range(2, ws_notic.max_row + 1):
            c_raw = str(ws_notic.cell(row=r, column=1).value or '').split('.')[0].strip()
            v_raw = ws_notic.cell(row=r, column=col_fecha_notic).value
            if c_raw:
                noticia_map[c_raw] = v_raw

    students = []
    asist_presentes = 0
    asist_ausentes = 0
    notic_presentes = 0
    notic_ausentes = 0
    pendientes_count = 0

    for r in range(2, ws_asist.max_row + 1):
        cuenta = ws_asist.cell(row=r, column=1).value
        nombre = ws_asist.cell(row=r, column=2).value
        if not cuenta and not nombre:
            continue

        cuenta_str = str(cuenta).split('.')[0].strip() if cuenta is not None else ''
        nombre_str = str(nombre).strip() if nombre is not None else ''

        # Asistencia
        val_asist = None
        status_asist = 'pending'
        if col_fecha_asist != -1:
            raw_v = ws_asist.cell(row=r, column=col_fecha_asist).value
            if raw_v is not None and str(raw_v).strip() != '':
                try:
                    val_num = int(float(raw_v))
                    val_asist = val_num
                    if val_num == 1:
                        status_asist = 'present'
                        asist_presentes += 1
                    elif val_num == 0:
                        status_asist = 'absent'
                        asist_ausentes += 1
                except:
                    pass

        # Noticia
        val_notic = None
        status_notic = 'pending'
        if cuenta_str in noticia_map:
            raw_n = noticia_map[cuenta_str]
            if raw_n is not None and str(raw_n).strip() != '':
                try:
                    val_n_num = int(float(raw_n))
                    val_notic = val_n_num
                    if val_n_num == 1:
                        status_notic = 'present'
                        notic_presentes += 1
                    elif val_n_num == 0:
                        status_notic = 'absent'
                        notic_ausentes += 1
                except:
                    pass

        if status_asist == 'pending' and status_notic == 'pending':
            pendientes_count += 1

        students.append({
            'row': r,
            'cuenta': cuenta_str,
            'nombre': nombre_str,
            'status_asistencia': status_asist,
            'val_asistencia': val_asist,
            'status_noticia': status_notic,
            'val_noticia': val_notic,
            # Compatibilidad
            'status': status_asist,
            'val_hoy': val_asist
        })

    wb.close()

    return jsonify({
        'ok': True,
        'materia': MATERIAS[materia_key],
        'fecha': fecha_filtro,
        'col_fecha_encontrada': col_fecha_asist != -1,
        'fechas_disponibles': fechas_clase,
        'total_alumnos': len(students),
        'asistencia_presentes': asist_presentes,
        'asistencia_ausentes': asist_ausentes,
        'noticia_presentes': notic_presentes,
        'noticia_ausentes': notic_ausentes,
        'pendientes': pendientes_count,
        # Compatibilidad retroactiva
        'presentes': asist_presentes,
        'ausentes': asist_ausentes,
        'students': students
    })


@app.route('/api/scan', methods=['POST'])
def registrar_escaneo_qr():
    """
    Procesa un escaneo de QR (o ID/nombre).
    Detecta si el QR declara específicamente si es de 'ASISTENCIA' o de 'NOTICIA',
    identifica al alumno (cuenta y nombre), marca 1 en la fecha indicada en la hoja
    correspondiente (Hoja2 para asistencia, Noticia para noticia) y sincroniza con Cloudflare R2.
    """
    data = request.get_json() or {}
    materia_key = data.get('materia', 'optimizacion')
    qr_data = str(data.get('qr_data', '')).strip()
    fecha = str(data.get('fecha', get_fecha_hoy_str())).strip()
    modo_solicitado = data.get('modo', 'asistencia').lower().strip()

    if not qr_data:
        return jsonify({'ok': False, 'error': 'Datos de QR vacíos.'}), 400

    path, filename = get_filepath(materia_key)
    if not path or not os.path.exists(path):
        return jsonify({'ok': False, 'error': 'Archivo no encontrado.'}), 404

    # Analizar si el QR declara el tipo: "ASISTENCIA|cuenta|nombre" o "NOTICIA|cuenta|nombre"
    partes = [p.strip() for p in qr_data.split('|')]
    tipo_declarado = None
    cuenta_buscada = ''
    nombre_buscado = ''

    if partes and partes[0].upper() in ['ASISTENCIA', 'NOTICIA']:
        tipo_declarado = partes[0].lower()
        cuenta_buscada = partes[1] if len(partes) > 1 else ''
        nombre_buscado = partes[2] if len(partes) > 2 else ''
    elif len(partes) > 2 and partes[2].upper() in ['ASISTENCIA', 'NOTICIA']:
        tipo_declarado = partes[2].lower()
        cuenta_buscada = partes[0]
        nombre_buscado = partes[1]
    else:
        # Formato clásico o fallback: cuenta|nombre
        cuenta_buscada = partes[0] if len(partes) > 0 else ''
        nombre_buscado = partes[1] if len(partes) > 1 else ''

    # El modo de registro se determina por el tipo declarado en el QR si existe, o el modo actual
    modo_final = tipo_declarado if tipo_declarado else modo_solicitado

    wb = openpyxl.load_workbook(path)
    ws_target = get_target_sheet(wb, modo_final)

    # Garantizar que la columna de la fecha exista aunque sea un día fuera del rol
    col_fecha = obtener_o_crear_columna_fecha(ws_target, fecha)

    alumno_encontrado = None
    row_encontrada = -1

    for r in range(2, ws_target.max_row + 1):
        c_val = str(ws_target.cell(row=r, column=1).value or '').split('.')[0].strip()
        n_val = str(ws_target.cell(row=r, column=2).value or '').strip()

        if cuenta_buscada and c_val == cuenta_buscada:
            alumno_encontrado = {'cuenta': c_val, 'nombre': n_val}
            row_encontrada = r
            break
        elif nombre_buscado and nombre_buscado.lower() in n_val.lower():
            alumno_encontrado = {'cuenta': c_val, 'nombre': n_val}
            row_encontrada = r
            break
        elif (not cuenta_buscada and not nombre_buscado) and (qr_data.lower() in n_val.lower() or qr_data == c_val):
            alumno_encontrado = {'cuenta': c_val, 'nombre': n_val}
            row_encontrada = r
            break

    if not alumno_encontrado:
        wb.close()
        return jsonify({'ok': False, 'error': f"Alumno no encontrado en la lista ({cuenta_buscada or nombre_buscado or qr_data})"}), 404

    # Escribir 1 (Presente / Entregó)
    cell = ws_target.cell(row=row_encontrada, column=col_fecha, value=1)
    cell.fill = FILL_GREEN
    cell.font = FONT_GREEN
    cell.alignment = ALIGN_CENTER

    ok, err = guardar_workbook_seguro(wb, path, filename)
    wb.close()

    if not ok:
        return jsonify({'ok': False, 'error': err}), 500

    nombre_hoja = "Noticia" if modo_final == "noticia" else "Asistencia"
    return jsonify({
        'ok': True,
        'modo': modo_final,
        'tipo_qr': tipo_declarado,
        'mensaje': f"Registrado para {alumno_encontrado['nombre']} en {nombre_hoja}",
        'alumno': alumno_encontrado
    })


@app.route('/api/marcar', methods=['POST'])
def marcar_asistencia_manual():
    """
    Marca un estado manual (1, 0, o vacío) para un alumno en una fecha dada en la hoja activa
    y sincroniza con Cloudflare R2. Si la columna de la fecha no existe, la crea.
    """
    data = request.get_json() or {}
    materia_key = data.get('materia', 'optimizacion')
    cuenta = str(data.get('cuenta', '')).strip()
    fecha = str(data.get('fecha', get_fecha_hoy_str())).strip()
    valor = data.get('valor')
    modo = data.get('modo', 'asistencia').lower().strip()

    path, filename = get_filepath(materia_key)
    if not path or not os.path.exists(path):
        return jsonify({'ok': False, 'error': 'Archivo no encontrado.'}), 404

    wb = openpyxl.load_workbook(path)
    ws_target = get_target_sheet(wb, modo)

    # Asegura o crea la columna de la fecha
    col_fecha = obtener_o_crear_columna_fecha(ws_target, fecha)

    target_row = -1
    for r in range(2, ws_target.max_row + 1):
        c_val = str(ws_target.cell(row=r, column=1).value or '').split('.')[0].strip()
        if c_val == cuenta:
            target_row = r
            break

    if target_row == -1:
        wb.close()
        return jsonify({'ok': False, 'error': 'Alumno no encontrado.'}), 404

    cell = ws_target.cell(row=target_row, column=col_fecha)
    if valor == 1:
        cell.value = 1
        cell.fill = FILL_GREEN
        cell.font = FONT_GREEN
        cell.alignment = ALIGN_CENTER
    elif valor == 0:
        cell.value = 0
        cell.fill = FILL_RED
        cell.font = FONT_RED
        cell.alignment = ALIGN_CENTER
    else:
        cell.value = None
        cell.fill = PatternFill(fill_type=None)
        cell.font = Font(name='Calibri')

    ok, err = guardar_workbook_seguro(wb, path, filename)
    wb.close()

    if not ok:
        return jsonify({'ok': False, 'error': err}), 500

    return jsonify({'ok': True, 'cuenta': cuenta, 'valor': valor})


@app.route('/api/finalizar', methods=['POST'])
def finalizar_clase():
    """
    Rellena todas las celdas vacías del día con 0 (Ausente) en la hoja activa
    y sincroniza con Cloudflare R2. Si la columna de la fecha no existe, la crea.
    """
    data = request.get_json() or {}
    materia_key = data.get('materia', 'optimizacion')
    fecha = str(data.get('fecha', get_fecha_hoy_str())).strip()
    modo = data.get('modo', 'asistencia').lower().strip()

    path, filename = get_filepath(materia_key)
    if not path or not os.path.exists(path):
        return jsonify({'ok': False, 'error': 'Archivo no encontrado.'}), 404

    wb = openpyxl.load_workbook(path)
    ws_target = get_target_sheet(wb, modo)

    # Asegura o crea la columna de la fecha
    col_fecha = obtener_o_crear_columna_fecha(ws_target, fecha)

    presentes = 0
    ausentes_rellenados = 0
    total = 0

    for r in range(2, ws_target.max_row + 1):
        cuenta_v = ws_target.cell(row=r, column=1).value
        nombre_v = ws_target.cell(row=r, column=2).value
        if not cuenta_v and not nombre_v:
            continue
        total += 1

        cell = ws_target.cell(row=r, column=col_fecha)
        raw = cell.value

        if raw is not None and str(raw).strip() != '':
            try:
                if int(float(raw)) == 1:
                    presentes += 1
            except:
                pass
        else:
            cell.value = 0
            cell.fill = FILL_RED
            cell.font = FONT_RED
            cell.alignment = ALIGN_CENTER
            ausentes_rellenados += 1

    ok, err = guardar_workbook_seguro(wb, path, filename)
    wb.close()

    if not ok:
        return jsonify({'ok': False, 'error': err}), 500

    nombre_hoja = "Noticia" if modo == "noticia" else "Asistencia (Hoja 2)"
    return jsonify({
        'ok': True,
        'fecha': fecha,
        'total': total,
        'presentes': presentes,
        'ausentes': ausentes_rellenados,
        'mensaje': f"Finalizado en {nombre_hoja}: {presentes} con valor 1 y {ausentes_rellenados} marcadas con 0."
    })


@app.route('/api/noticias', methods=['POST'])
def agregar_noticia():
    """Agrega una nueva noticia/aviso a la hoja Noticias del Excel y sincroniza con R2."""
    data = request.get_json() or {}
    materia_key = data.get('materia', 'optimizacion')
    titulo = str(data.get('titulo', '')).strip()
    cuerpo = str(data.get('cuerpo', '')).strip()
    prioridad = str(data.get('prioridad', 'Informativo')).strip()
    fecha = str(data.get('fecha', get_fecha_hoy_str())).strip()

    if not titulo or not cuerpo:
        return jsonify({'ok': False, 'error': 'Título y contenido son requeridos.'}), 400

    path, filename = get_filepath(materia_key)
    if not path or not os.path.exists(path):
        return jsonify({'ok': False, 'error': 'Archivo no encontrado.'}), 404

    wb = openpyxl.load_workbook(path)
    if 'Noticias' not in wb.sheetnames:
        ws_not = wb.create_sheet('Noticias')
        for c_i, h in enumerate(['FECHA', 'TÍTULO', 'PRIORIDAD', 'DESCRIPCIÓN / AVISO'], start=1):
            c = ws_not.cell(row=1, column=c_i, value=h)
            c.fill = FILL_HDR
            c.font = FONT_HDR
    else:
        ws_not = wb['Noticias']

    next_row = ws_not.max_row + 1
    ws_not.cell(row=next_row, column=1, value=fecha)
    ws_not.cell(row=next_row, column=2, value=titulo)
    ws_not.cell(row=next_row, column=3, value=prioridad)
    ws_not.cell(row=next_row, column=4, value=cuerpo)

    ok, err = guardar_workbook_seguro(wb, path, filename)
    wb.close()

    if not ok:
        return jsonify({'ok': False, 'error': err}), 500

    return jsonify({'ok': True, 'mensaje': 'Noticia agregada con éxito en el archivo Excel.'})


@app.route('/archivos_excel/<path:filename>')
def serve_excel_file(filename):
    """
    Permite descargar los archivos Excel directamente.
    Asegura que se entregue la versión más reciente descargada de R2.
    """
    local_path = r2_storage.ensure_file_local(filename)
    if not os.path.exists(local_path):
        return jsonify({'ok': False, 'error': 'Archivo no encontrado'}), 404
    return send_file(local_path, as_attachment=True, download_name=filename)


# ─── INICIO DEL SERVIDOR ──────────────────────────────────────────────────────

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print("=" * 60)
    print("  📋 PORTAL ACADÉMICO — CLOUDFLARE R2 + RENDER")
    print("=" * 60)
    print(f"  ☁️ Cloudflare R2 : {'Conectado' if r2_storage.is_r2_configured() else 'Modo Local /tmp'}")
    print(f"  📅 Fecha actual  : {get_fecha_hoy_str()}")
    print(f"  🌐 Dirección Web : http://localhost:{port}")
    print("=" * 60)
    print("  Presiona Ctrl+C en esta consola para detener el servidor.\n")

    app.run(host='0.0.0.0', port=port, debug=False)
